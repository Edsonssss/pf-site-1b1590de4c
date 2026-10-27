#!/usr/bin/env python3
"""Пересборка Excel-таблицы из portfolio_data.json (авто-режим GitHub Actions)."""
import json, os, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
d = json.load(open(os.path.join(BASE, 'portfolio_data.json'), encoding='utf-8'))
stocks, funds, cash, sum_, m = d['stocks'], d['funds'], d['cash'], d['summary'], d['macro']
tot = sum_['total_value']
dt = datetime.date.fromisoformat(d['meta']['date'])
DATE = dt.strftime('%d.%m.%Y')
pd_ = sorted({s.get('prev_date') for s in stocks if s.get('prev_date')})[-1]
PREV = f"{pd_[8:10]}.{pd_[5:7]}.{pd_[:4]}"
OUT = os.path.join(BASE, 'Портфель_и_прогнозы_ТБанк_08.10.2026.xlsx')

wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Портфель и прогнозы'
HDR = ['№','Категория','Тикер','Название','Кол-во',f'Текущая цена, ₽ ({DATE})',f'Цена за пред. торг. день, ₽ ({PREV})','Изменение, %','Текущая стоимость, ₽','Доля, % *','Рекомендация Т-Банк','Целевая цена, ₽','Потенциал, %','Диапазон целей, ₽','Ближайший дивиденд, ₽/акц.','Дата отсечки','Див. доходность за 12 мес','Статус дивиденда','Примечания / прогноз']
W = [4,10,9,28,10,14,16,11,14,9,14,12,10,14,15,26,11,17,50]
NC = len(HDR)
thin = Side(style='thin', color='BFBFBF'); border = Border(left=thin, right=thin, top=thin, bottom=thin)
hdr_fill = PatternFill('solid', fgColor='1F3864'); hdr_font = Font(bold=True, color='FFFFFF', size=10)
buy_fill = PatternFill('solid', fgColor='C6EFCE'); buy_font = Font(color='006100', bold=True, size=10)
hold_fill = PatternFill('solid', fgColor='FFEB9C'); hold_font = Font(color='9C6500', bold=True, size=10)
none_fill = PatternFill('solid', fgColor='D9D9D9'); none_font = Font(color='595959', size=10)
div_fill = PatternFill('solid', fgColor='E2EFDA'); div_font = Font(color='006100', size=10, bold=True)
fcst_font = Font(color='1F5C99', size=9); tot_fill = PatternFill('solid', fgColor='DDEBF7')
sec_fill = PatternFill('solid', fgColor='D9D9D9'); grand_fill = PatternFill('solid', fgColor='BDD7EE')
pos_font = Font(color='006100', size=10); negf = Font(color='9C0006', size=10)
def price_fmt(v):
    if v is None: return 'General'
    if v < 1: return '0.00000'
    return '#,##0.00'

ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=NC)
c = ws.cell(row=1, column=1, value=f'Инвестиционный портфель — сводные прогнозы Т-Банка (Т-Инвестиции) и дивиденды на {DATE} (автообновление)')
c.font = Font(bold=True, size=13, color='1F3864'); c.alignment = Alignment(horizontal='center')
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=NC)
c = ws.cell(row=2, column=1, value=f'Источники: сводный прогноз Т-Банка — на 08.10.2026; цены — MOEX ISS ({DATE}, закрытие {PREV}); паи — investfunds.ru; дивиденды — tbank.ru + smart-lab.ru. * Доли: в строках акций — внутри акции-части, фондов — внутри фондов, в итогах и кэше — в портфеле.')
c.font = Font(italic=True, size=9, color='808080'); c.alignment = Alignment(horizontal='center', wrap_text=True)
ws.row_dimensions[2].height = 30
hr = 3
for i, (hh, w) in enumerate(zip(HDR, W), 1):
    cell = ws.cell(row=hr, column=i, value=hh); cell.fill = hdr_fill; cell.font = hdr_font
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); cell.border = border
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[hr].height = 48; ws.freeze_panes = 'A4'

def write_row(rw, s, is_fund=False):
    if is_fund:
        nd, dst, st, y12 = '— (выплат нет)', '—', '—', None
        name, cat, tick = s['name'], 'Фонды', '—'
        qty, cur, val = s['qty'], s['current_nav'], s['value']
        prev, day = None, None
        rec, tgt, pot, rng = '—', None, None, None
        note = s['forecast']
    else:
        nd, dst, st = s['dividend_next'], s['dividend_date'], s['dividend_status']
        y12 = (s['div_yield_12m_pct'] / 100) if s['div_yield_12m_pct'] else None
        name, cat, tick = s['name'], 'Акции', s['ticker']
        qty, cur, val = s['qty'], s['current_price'], s['value']
        prev = s.get('prev_close'); day = (s.get('day_change_pct') / 100 if s.get('day_change_pct') is not None else None)
        rec = s['recommendation'] or 'нет прогноза'; tgt = s['target_price']
        pot = (s['upside_pct'] / 100 if s['upside_pct'] is not None else None)
        rng = f"{s['range_min']:g} — {s['range_max']:g}" if s['range_min'] is not None else None
        note = s['note']
    den = sum_['stocks_value'] if cat == 'Акции' else (sum_['funds_value'] if cat == 'Фонды' else tot)
    vals = [s['n'], cat, tick, name, qty, cur, prev, day, val, val / den, rec, tgt, pot, rng, nd, dst, y12, st, note]
    for i, v in enumerate(vals, 1):
        cell = ws.cell(row=rw, column=i, value=v); cell.border = border; cell.font = Font(size=10)
        if i == 1: cell.alignment = Alignment(horizontal='center')
        if i == 5: cell.number_format = '#,##0.###'
        if i in (6, 7, 12): cell.number_format = price_fmt(v if isinstance(v, (int, float)) else None)
        if i == 8 and isinstance(v, float): cell.number_format = '+0.00%;-0.00%'; cell.font = pos_font if v >= 0 else negf
        if i == 9: cell.number_format = '#,##0'
        if i == 10: cell.number_format = '0.00%'
        if i == 13 and isinstance(v, float): cell.number_format = '+0.0%;-0.0%'; cell.font = pos_font if v >= 0 else negf
        if i == 11:
            cell.alignment = Alignment(horizontal='center')
            if v == 'Покупать': cell.fill = buy_fill; cell.font = buy_font
            elif v == 'Держать': cell.fill = hold_fill; cell.font = hold_font
            elif v == 'нет прогноза': cell.fill = none_fill; cell.font = none_font
        if i == 15 and isinstance(v, (int, float)): cell.number_format = '#,##0.00'; cell.fill = div_fill; cell.font = div_font
        if i == 17 and isinstance(v, float): cell.number_format = '0.0%'
        if i in (15, 16, 18) and isinstance(v, str) and v != '—':
            cell.font = fcst_font
            if st == 'утвержден СД' and i in (15, 16): cell.font = div_font
        if i in (14, 16, 18, 19): cell.alignment = Alignment(wrap_text=True, vertical='top')
        if i == 19: cell.font = Font(size=9)
    return rw + 1

def total_row(rw, label, val, share, extra=None, fill=tot_fill):
    ws.merge_cells(start_row=rw, start_column=1, end_row=rw, end_column=8)
    c = ws.cell(row=rw, column=1, value=label); c.font = Font(bold=True, size=10); c.fill = fill; c.alignment = Alignment(horizontal='right')
    c2 = ws.cell(row=rw, column=9, value=val); c2.number_format = '#,##0'; c2.font = Font(bold=True, size=10); c2.fill = fill
    c3 = ws.cell(row=rw, column=10, value=share); c3.number_format = '0.00%'; c3.font = Font(bold=True, size=10); c3.fill = fill
    for col in range(11, NC + 1): ws.cell(row=rw, column=col).fill = fill
    if extra:
        ce = ws.cell(row=rw, column=11, value=extra); ce.font = Font(bold=True, size=9); ce.fill = fill; ce.alignment = Alignment(horizontal='left')
    for col in range(1, NC + 1): ws.cell(row=rw, column=col).border = border
    return rw + 1

rw = hr + 1
for s in stocks: rw = write_row(rw, s)
rw = total_row(rw, f"ИТОГО АКЦИИ ({len(stocks)} позиций)", sum_['stocks_value'], sum_['stocks_value'] / tot,
               extra=f"Срвзв. потенциал: +{sum_['weighted_upside_pct']}% · Покупать {sum_['recommendations']['buy']} / Держать {sum_['recommendations']['hold']} / без прогноза {sum_['recommendations']['no_forecast']}")
for f in funds: rw = write_row(rw, f, is_fund=True)
rw = total_row(rw, f"ИТОГО ФОНДЫ ({len(funds)} позиции)", sum_['funds_value'], sum_['funds_value'] / tot)
vals = [cash['n'], 'Cash', '—', cash['name'], None, None, None, None, cash['value'], cash['value'] / tot, '—', None, None, None, '—', '—', None, '—', cash['note']]
for i, v in enumerate(vals, 1):
    cell = ws.cell(row=rw, column=i, value=v); cell.border = border; cell.font = Font(size=10)
    if i == 9: cell.number_format = '#,##0'
    if i == 10: cell.number_format = '0.00%'
    if i == 11: cell.alignment = Alignment(horizontal='center')
    if i == 19: cell.alignment = Alignment(wrap_text=True, vertical='top'); cell.font = Font(size=9)
rw += 1
rw = total_row(rw, 'ИТОГО CASH', cash['value'], cash['value'] / tot)
rw = total_row(rw, 'ИТОГО ПОРТФЕЛЬ (ликвидная часть)', tot, 1.0, fill=grand_fill)
ce = ws.cell(row=rw - 1, column=15, value=(f"Ближайшие дивиденды (утв.): {sum_['upcoming_dividends_gross']:,.0f} ₽".replace(',', ' ') if sum_['upcoming_dividends_gross'] else 'Ближайших утверждённых отсечек нет'))
ce.font = Font(bold=True, size=9, color='006100'); ce.alignment = Alignment(horizontal='left')
fam = d.get('family') or {}
if fam:
    vals = ['—', 'Cash (жена)', '—', fam['name'], None, None, None, None, fam['value'], None, '—', None, None, None, '—', '—', None, '—', fam.get('note', '')]
    for i, v in enumerate(vals, 1):
        cell = ws.cell(row=rw, column=i, value=v); cell.border = border; cell.font = Font(size=10)
        if i == 9: cell.number_format = '#,##0'
        if i == 19: cell.alignment = Alignment(wrap_text=True, vertical='top'); cell.font = Font(size=9)
    rw += 1
    rw = total_row(rw, 'ВСЕГО СЕМЕЙНЫЙ ПОРТФЕЛЬ (ликвидный)', sum_.get('total_family', round(tot + fam['value'], 2)), None,
                   extra=f"ваш {tot:,.0f} ₽ + жены {fam['value']:,.0f} ₽".replace(',', ' '), fill=grand_fill)
if d.get('blocked'):
    ws.merge_cells(start_row=rw, start_column=1, end_row=rw, end_column=NC)
    c = ws.cell(row=rw, column=1, value='ЗАБЛОКИРОВАННЫЕ АКТИВЫ — справочно, вне ликвидного итога')
    c.font = Font(bold=True, size=10, color='404040'); c.fill = sec_fill; c.alignment = Alignment(horizontal='left')
    for col in range(1, NC + 1): ws.cell(row=rw, column=col).border = border; ws.cell(row=rw, column=col).fill = sec_fill
    rw += 1
    for b0 in d['blocked']:
        px = round(b0['price_usd'] * m['usd_rate_cbr'], 2) if b0.get('price_usd') else None
        vals = [b0['n'], 'Заблокировано', b0['ticker'], b0['name'], b0['qty'], px, None, None, b0['value_rub'], None, '—', None, None, None, '—', '—', None, 'заблокировано', b0['note']]
        for i, v in enumerate(vals, 1):
            cell = ws.cell(row=rw, column=i, value=v); cell.border = border; cell.font = Font(size=10)
            if i == 1: cell.alignment = Alignment(horizontal='center')
            if i == 6 and v: cell.number_format = '#,##0.00'
            if i == 9: cell.number_format = '#,##0'
            if i == 11: cell.alignment = Alignment(horizontal='center')
            if i in (18, 19): cell.alignment = Alignment(wrap_text=True, vertical='top'); cell.font = Font(size=9)
        rw += 1
    ws.merge_cells(start_row=rw, start_column=1, end_row=rw, end_column=8)
    c = ws.cell(row=rw, column=1, value='ИТОГО С УЧЁТОМ ЗАБЛОКИРОВАННЫХ'); c.font = Font(bold=True, size=10); c.fill = grand_fill; c.alignment = Alignment(horizontal='right')
    c2 = ws.cell(row=rw, column=9, value=sum_['total_including_blocked']); c2.number_format = '#,##0'; c2.font = Font(bold=True, size=10); c2.fill = grand_fill
    c3 = ws.cell(row=rw, column=10, value=f"ликвидная часть: {tot:,.0f} ₽ + заблокировано {sum_['blocked_value_rub']:,.0f} ₽".replace(',', ' ')); c3.font = Font(bold=True, size=9); c3.fill = grand_fill; c3.alignment = Alignment(horizontal='left')
    for col in range(1, NC + 1): ws.cell(row=rw, column=col).border = border; ws.cell(row=rw, column=col).fill = grand_fill

ws2 = wb.create_sheet('Сводка и макропрогнозы')
ws2.column_dimensions['A'].width = 4; ws2.column_dimensions['B'].width = 118
def sec(rw_, t_):
    c = ws2.cell(row=rw_, column=2, value=t_); c.font = Font(bold=True, size=12, color='1F3864'); return rw_ + 1
def para(rw_, t_):
    c = ws2.cell(row=rw_, column=2, value=t_); c.font = Font(size=10); c.alignment = Alignment(wrap_text=True, vertical='top')
    ws2.row_dimensions[rw_].height = max(15, 15 * (len(t_) // 110 + 1)); return rw_ + 1
r = 1
r = sec(r, f'СВОДКА: ПРОГНОЗЫ Т-БАНКА (08.10.2026), ЦЕНЫ И СОСТАВ НА {DATE} (АВТООБНОВЛЕНИЕ)')
r = para(r, f"• Рекомендации: ПОКУПАТЬ — {sum_['recommendations']['buy']}; ДЕРЖАТЬ — {sum_['recommendations']['hold']}; ПРОДАВАТЬ — 0; без прогноза — {sum_['recommendations']['no_forecast']} (FEES, PMSB).")
r = para(r, '• Средневзвешенный потенциал роста акций к целевым ценам: +%.1f%% (без учёта дивидендов).' % sum_['weighted_upside_pct'])
r += 1
r = sec(r, 'ДИВИДЕНДЫ')
if d['upcoming_dividends']:
    for u in d['upcoming_dividends']:
        rd = u['record_date']; lb = u['last_buy_date']
        r = para(r, f"• {u['ticker']}: {u['dividend']} ₽/акц., отсечка {rd[8:10]}.{rd[5:7]}.{rd[:4]} (купить до {lb[8:10]}.{lb[5:7]}.{lb[:4]}), {u['qty']} акций → {u['gross']:,.0f} ₽.".replace(',', ' '))
else:
    r = para(r, '• Ближайших утверждённых отсечек нет. Выплаты T, GMKN (12.10) и SIBN (13.10) зафиксированы и придут в конце октября — начале ноября.')
r = para(r, f"• Дивиденды за 12 месяцев (факт): ≈ {sum_['dividends_12m_gross']:,.0f} ₽ — {sum_['dividends_12m_yield_pct']}% годовых на портфель.".replace(',', ' '))
r += 1
r = sec(r, 'МАКРО И СТРАТЕГИЯ')
r = para(r, '• IMOEX: 2 300–2 500 п. к концу 2026; 2 600–2 900 п. за 12 мес (Т-Инвестиции, 01.08.2026). Ключевая ставка: 14,00% (11.09.2026).')
r += 1
r = sec(r, 'ФОНДЫ')
for f in funds:
    p_ = f['performance_pct']
    r = para(r, f"• {f['name']}: пай {f['current_nav']:,.2f} ₽ ({f.get('nav_date', '—')}); 1 мес {p_['m1']:+}%, 1 год {p_['y1']:+}%. {f['forecast']}".replace(',', ' '))
fam2 = d.get('family') or {}
if fam2:
    r += 1
    r = sec(r, 'СЕМЕЙНЫЙ БЛОК (вне расчётов основного портфеля)')
    r = para(r, f"• {fam2['name']}: {fam2['value']:,.0f} ₽. {fam2.get('note', '')}".replace(',', ' '))
    r = para(r, f"• Семейный итог (ликвидный): {sum_.get('total_family', sum_['total_value'] + fam2['value']):,.0f} ₽ = ваш {sum_['total_value']:,.0f} ₽ + {fam2['value']:,.0f} ₽.".replace(',', ' '))
if d.get('blocked'):
    r += 1
    r = sec(r, 'ЗАБЛОКИРОВАННЫЕ АКТИВЫ')
    for b0 in d['blocked']:
        cur = f"{b0['qty']} шт × {b0['price_usd']} $ ≈ {b0['value_rub']:,.0f} ₽" if b0.get('price_usd') else f"{b0['qty']} шт, оценка {b0['value_rub']:,.0f} ₽"
        r = para(r, f"• {b0['name']}: {cur}. {b0['note']}")
    r = para(r, f"• Итого с учётом заблокированных: {sum_['total_including_blocked']:,.0f} ₽ (ликвидная часть {tot:,.0f} ₽).".replace(',', ' '))
r += 1
r = sec(r, 'ИСТОЧНИКИ')
r = para(r, f'• tbank.ru/invest (08.10.2026); MOEX ISS ({DATE}, закрытие {PREV}); investfunds.ru; smart-lab.ru/dividends; cbr.ru. Состав и количества обновляются вручную.')
r = para(r, '⚠️ Материал информационный, не является индивидуальной инвестиционной рекомендацией.')

# ---------- лист «Остатки по дням» (ежедневный учёт + график) ----------
hist_path = os.path.join(BASE, 'history.json')
if os.path.exists(hist_path):
    from openpyxl.chart import LineChart, Reference
    from openpyxl.chart.marker import Marker
    hist = json.load(open(hist_path, encoding='utf-8'))
    days = hist.get('days', [])
    ws3 = wb.create_sheet('Остатки по дням', 1)
    ws3.merge_cells('A1:J1')
    c = ws3.cell(row=1, column=1, value=f'Учёт ежедневных остатков — Акции, Фонды, Счета, Сэйв жены (обновлено: {hist.get("updated", DATE)})')
    c.font = Font(bold=True, size=13, color='1F3864'); c.alignment = Alignment(horizontal='center')
    ws3.merge_cells('A2:J2')
    c = ws3.cell(row=2, column=1, value=hist.get('note', ''))
    c.font = Font(italic=True, size=9, color='808080'); c.alignment = Alignment(wrap_text=True, vertical='top')
    ws3.row_dimensions[2].height = 42
    HDR3 = ['Дата', 'Сэйв жены, ₽', 'Тинькоф — акции, ₽', 'IH+БКС — фонды, ₽', 'Счета (кэш), ₽',
            'ИТОГО портфель, ₽', 'Заблокировано, ₽', 'Итого с заблокир., ₽', 'Семейный итог, ₽',
            'Примечание / источник']
    W3 = [12, 14, 15, 16, 14, 17, 15, 17, 16, 66]
    hr3 = 4
    for i, (hh, w) in enumerate(zip(HDR3, W3), 1):
        cell = ws3.cell(row=hr3, column=i, value=hh); cell.fill = hdr_fill; cell.font = hdr_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); cell.border = border
        ws3.column_dimensions[get_column_letter(i)].width = w
    ws3.row_dimensions[hr3].height = 30
    ws3.freeze_panes = f'A{hr3 + 1}'
    for k, dy in enumerate(days):
        r_ = hr3 + 1 + k
        dd = dy['date']
        vals = [f"{dd[8:10]}.{dd[5:7]}.{dd[:4]}", dy.get('save_wife'), dy.get('stocks'), dy.get('funds'),
                dy.get('cash'), dy.get('total'), dy.get('blocked_rub'),
                dy.get('total_incl_blocked'), dy.get('total_family'), dy.get('source', '')]
        for i, v in enumerate(vals, 1):
            cell = ws3.cell(row=r_, column=i, value=v); cell.border = border; cell.font = Font(size=10)
            if i == 1: cell.alignment = Alignment(horizontal='center')
            elif 2 <= i <= 9: cell.number_format = '#,##0'
            else: cell.alignment = Alignment(wrap_text=True, vertical='top'); cell.font = Font(size=9, color='595959')
        if k % 2 == 1:
            for i in range(1, 11): ws3.cell(row=r_, column=i).fill = PatternFill('solid', fgColor='FAFBFD')
        for i in (6, 9): ws3.cell(row=r_, column=i).font = Font(size=10, bold=True)
    if days:
        chart = LineChart()
        chart.title = 'Динамика остатков: помесячно с 03.2023 (на 1-е число) + ежедневно, ₽'
        chart.style = 2
        chart.y_axis.title = '₽'; chart.x_axis.title = 'Дата'
        chart.y_axis.numFmt = '#,##0'
        chart.y_axis.delete = False; chart.x_axis.delete = False
        chart.width = 30; chart.height = 14
        chart.legend.position = 'b'
        chart.displayBlanksAs = 'gap'
        COLS = {2: ('A855F7', 19050), 3: ('2563EB', 19050), 4: ('16A34A', 19050),
                5: ('F59E0B', 19050), 6: ('111827', 28575), 9: ('DC2626', 28575)}
        for col, (clr, wd) in COLS.items():
            ref = Reference(ws3, min_col=col, min_row=hr3, max_row=hr3 + len(days))
            chart.add_data(ref, titles_from_data=True)
        chart.set_categories(Reference(ws3, min_col=1, min_row=hr3 + 1, max_row=hr3 + len(days)))
        msize = 5 if len(days) <= 20 else 3
        for s_ in chart.series:
            s_.smooth = False
            s_.graphicalProperties.line.solidFill = '808080'
            s_.marker = Marker(symbol='circle', size=msize)
        for (col, (clr, wd)), s_ in zip(sorted(COLS.items()), chart.series):
            s_.graphicalProperties.line.solidFill = clr
            s_.graphicalProperties.line.width = wd
            s_.marker.graphicalProperties.solidFill = clr
        if len(days) >= 2:
            dash = chart.series[-1].graphicalProperties.line
            dash.dashStyle = 'dash'
        ws3.add_chart(chart, f'A{hr3 + len(days) + 3}')

wb.save(OUT)

# xlsx в Google Drive (если задан ключ)
key_raw = os.environ.get('DRIVE_SA_KEY', '')
if key_raw:
    try:
        import urllib.request, ssl, jwt as _jwt, datetime as _dt, urllib.parse
        CTX = ssl.create_default_context(); CTX.check_hostname = False; CTX.verify_mode = ssl.CERT_NONE
        KEY = json.loads(key_raw); FOLDER = '0B-tnKJ9INY9WeVRsdndmWlM1Nzg'
        now = int(_dt.datetime.now().timestamp())
        claims = {'iss': KEY['client_email'], 'scope': 'https://www.googleapis.com/auth/drive', 'aud': KEY['token_uri'], 'exp': now + 3600, 'iat': now}
        t = _jwt.encode(claims, KEY['private_key'], algorithm='RS256')
        resp = json.loads(urllib.request.urlopen(urllib.request.Request(KEY['token_uri'], data=('grant_type=urn%3Aietf%3Aparams%3Aoauth%3Agrant-type%3Ajwt-bearer&assertion=' + t).encode(), headers={'Content-Type': 'application/x-www-form-urlencoded'}), timeout=30, context=CTX).read())
        AT = resp['access_token']
        nm = os.path.basename(OUT)
        q = urllib.parse.quote(f"name = '{nm}' and '{FOLDER}' in parents and trashed = false")
        rr = json.loads(urllib.request.urlopen(urllib.request.Request(f'https://www.googleapis.com/drive/v3/files?q={q}&fields=files(id)', headers={'Authorization': 'Bearer ' + AT}), timeout=30, context=CTX).read())
        if rr['files']:
            urllib.request.urlopen(urllib.request.Request(f"https://www.googleapis.com/upload/drive/v3/files/{rr['files'][0]['id']}?uploadType=media", data=open(OUT, 'rb').read(), method='PATCH', headers={'Authorization': 'Bearer ' + AT, 'Content-Type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'}), timeout=60, context=CTX)
            print('Drive: xlsx обновлён')
    except Exception as e:
        print('Drive xlsx error:', e)
print('Excel пересобран:', OUT)
